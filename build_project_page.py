from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pathlib import Path

root = Path.cwd()
out = root / 'Self_Driving_Car_Mini_Project.docx'
img = root / 'self_driving_car_reference.png'
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(0.42)
sec.bottom_margin = Inches(0.38)
sec.left_margin = Inches(0.62)
sec.right_margin = Inches(0.62)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.0

def para(text='', size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=0, line=1.0, italic=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    return p

para('AUTONOMOUS LINE FOLLOWING AND OBSTACLE AVOIDANCE CAR',18,True,WD_ALIGN_PARAGRAPH.CENTER,before=14,after=2)
para('Mini Project',12,True,WD_ALIGN_PARAGRAPH.CENTER,after=9)
para('Group Members',11,True,WD_ALIGN_PARAGRAPH.CENTER,after=3)
para('Registration Numbers: RA2611067010041, RA2611067010083, RA2611067010084, RA2611067010098, RA2611067010040',10.2,False,WD_ALIGN_PARAGRAPH.CENTER,after=2,line=1.08)
para('Section M1',10.5,True,WD_ALIGN_PARAGRAPH.CENTER,after=9)
para('B.Tech in Electrical and Electronics Engineering',11,False,WD_ALIGN_PARAGRAPH.CENTER,after=1)
para('VLSI Design and Technology',10.5,False,WD_ALIGN_PARAGRAPH.CENTER,after=4)
para('College of Engineering & Technology',10.5,False,WD_ALIGN_PARAGRAPH.CENTER,after=1)
para('SRM Institute of Science and Technology, Kattankulathur, Chengalpattu District, Tamil Nadu',9.8,False,WD_ALIGN_PARAGRAPH.CENTER,after=8)
para('Faculty In-Charge',11,True,WD_ALIGN_PARAGRAPH.CENTER,after=1)
para('Radhamani A V',10.5,False,WD_ALIGN_PARAGRAPH.CENTER,after=1)
para('Department of Electrical and Electronics Engineering',9.8,False,WD_ALIGN_PARAGRAPH.CENTER,after=9)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(4)
p.paragraph_format.line_spacing = 1.05
r = p.add_run('Summary: ')
r.bold = True
body = ('This mini project presents a low-cost autonomous car that follows a marked path and responds to obstacles on its route. '
        'An infrared sensor array detects the line and helps the controller correct the car’s direction. An ultrasonic sensor measures the distance to objects ahead; when an obstacle is detected, the car stops, scans for a clear direction, and turns before searching for the line again. '
        'A microcontroller processes the sensor readings and sends control signals through a dual motor driver to the two geared motors. The project brings together embedded C/C++ programming, sensor interfacing, motor control, and basic autonomous navigation. '
        'The prototype will be demonstrated on a controlled indoor course using a high-contrast track and safe, stationary obstacles. Its performance can be evaluated by line-tracking accuracy, obstacle response, and whether it can rejoin the route after a detour.')
r = p.add_run(body)
r.font.name = 'Times New Roman'; r.font.size = Pt(10.5); r._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(3)
p.paragraph_format.space_after = Pt(2)
p.add_run().add_picture(str(img), width=Inches(6.45))
para('Figure. Concept image of the autonomous car with line-following and obstacle-detection sensors.',9.2,False,WD_ALIGN_PARAGRAPH.CENTER,after=0,italic=True)

doc.core_properties.title = 'Autonomous Line Following and Obstacle Avoidance Car'
doc.core_properties.subject = 'Mini Project Cover Page'
doc.save(out)
print(out)
